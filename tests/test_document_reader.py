"""Unit tests for multi-format (PDF/DOCX/TXT) document text extraction."""
import io

import docx
import pytest
from pypdf import PdfWriter

from src.ingestion.document_reader import MAX_FILE_SIZE_BYTES, extract_text


def test_extract_text_from_txt_bytes():
    result = extract_text("notes.txt", "Hello world".encode("utf-8"))
    assert result == "Hello world"


def test_extract_text_from_docx_bytes():
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Grant application paragraph one.")
    document.add_paragraph("Second paragraph.")
    document.save(buffer)

    result = extract_text("application.docx", buffer.getvalue())
    assert "Grant application paragraph one." in result
    assert "Second paragraph." in result


def test_extract_text_from_pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    # A blank page yields empty text but must not raise.
    result = extract_text("application.pdf", buffer.getvalue())
    assert result == ""


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(ValueError):
        extract_text("malware.exe", b"binary")


def test_extract_text_rejects_oversized_file():
    with pytest.raises(ValueError):
        extract_text("big.txt", b"0" * (MAX_FILE_SIZE_BYTES + 1))
